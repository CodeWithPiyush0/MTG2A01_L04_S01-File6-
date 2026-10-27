"""The lesson's voice-over manifest - lesson_vo_manifest.xlsx (the format of HI02H11_L02_S02's, File3).

Four sheets:
  VO Manifest     every clip where it is heard, in game order (start screen, transitions, pages 1-10, celebration):
                  page, screen, #, Audio ID, the Hindi script (spoken exactly), what it is, when it plays, the other
                  pages that play the same recording, duration, status
  Recording list  each recording once, in the order it is first heard, with its file name
  Sound effects   the sound files the lesson plays (supplied, not recorded)
  Recording spec  scope, folder, format, how to put new recordings in

Reads card.json (the build's own text and the page order) and the clips on disk (durations).
    python tools/gen_vo_manifest.py
"""
import json, os, subprocess
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CARD = json.load(open(os.path.join(ROOT, "card.json"), encoding="utf-8"))
TEXT = CARD["assets"]["audio_text"]
VO_DIR = os.path.join(ROOT, "assets", "voiceover")
SFX_DIR = os.path.join(ROOT, "assets", "sfx")
OUT = os.path.join(ROOT, "lesson_vo_manifest.xlsx")
TITLE = CARD["title"]["hi"] + " (" + CARD["skill_code"] + ")"

COPIED = set()          # (the transition lines were HI02H11's clips until the 2026-10-27 recordings)


def duration(path):
    try:
        return round(float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", path],
                                          capture_output=True, text=True).stdout.strip()), 2)
    except Exception:
        return None


# ---- what each role is and when it plays (the build's behaviour) ----
ROLE = {
    # T1
    ("T1", "a"): ("story - Meera's blocks", "as the page opens; Meera's picture and her 25 blocks come in, the card 25; her picture lights on «मीरा»"),
    ("T1", "b"): ("story - 13 given to Aman", "next; Aman's picture comes in, then − 13; the pictures light on «मीरा» / «अमन»"),
    ("T1", "ask"): ("story - the question", "next; = ? comes in and pulses; Meera's picture lights on «मीरा»"),
    ("T1", "learn"): ("story - what we will learn", "next; 25 and 13 light as they are said; then the next button"),
    # T2
    ("T2", "head"): ("demo - opening", "as the page opens; the 25 card pops"),
    ("T2", "intro"): ("demo - build 25", "next; the empty table and the + box come in"),
    ("T2", "digits"): ("demo - 25 = 2 tens 5 ones", "next; the digits and the + halves light / pulse on «2 दहाई» / «5 इकाई»"),
    ("T2", "tens"): ("demo - the tens", "Swiftee's hand taps the rod + twice"),
    ("T2", "ones"): ("demo - the ones", "Swiftee's hand taps the cube + five times"),
    ("T2", "built"): ("demo - 25 is built", "the table pops"),
    ("T2", "sub"): ("demo - take 13 away", "25 − 13 = ? comes in part by part with the words; 13's digits light on «1 दहाई» / «3 इकाई»"),
    ("T2", "rm_ones"): ("demo - ones first", "then Swiftee's hand takes 3 ones away"),
    ("T2", "num"): ("demo - the count", "one per tap of Swiftee's hand on the ones: «एक», «दो», «तीन»"),
    ("T2", "ones_left"): ("demo - ones left", "the ones column beats on «2 इकाई»"),
    ("T2", "rm_tens"): ("demo - the child takes the ten", "the child taps one ten (the hand after 3s)"),
    ("T2", "tens_left"): ("demo - tens left", "the tens column beats on «1 दहाई»"),
    ("T2", "total"): ("demo - the result", "the columns beat; = 12 comes in beside the table on «यानी 12»"),
    ("T2", "story_a"): ("story's answer", "T1's screen comes back; on «तेरह घटाए» 13 of Meera's blocks fly into Aman's picture (hers grey); on «बचे 12» her 12 beat"),
    ("T2", "story_b"): ("story's answer - the box", "the ? becomes 12 on «12 आएगा»; confetti"),
    ("T2", "story_c"): ("story's answer - the end", "the 12 beat, Meera's picture lights on «मीरा»; then the next button"),
    # the guided / practice pages
    "prompt": ("page instruction", "as the page opens (the speaker replays it)"),
    "build": ("build instruction", "after the instruction; the + box comes in, the rod + pulses on its «… दहाई», the cube + on «… इकाई»"),
    "sub": ("what is taken away", "after the build (or the opening): the number pops, its digits light on «… दहाई» / «… इकाई»; on «इसमें» the whole table lights"),
    "remove": ("take-away instruction", "next; the ones column beats on its «… इकाई», the tens on «… दहाई»; then «अब इकाई हटाइए।»"),
    "rm_ones": ("take-away step - the ones", "next; the ones column beats on its «… इकाई»"),
    "zero": ("zero ones", "straight after the ones step: no one is tappable"),
    "rm_tens": ("take-away step - the tens", "when the ones are done (at once when there are none to take)"),
    "ask": ("the answer question", "when the take-away is done; the cards / the keypad come in (said again after 7s idle)"),
    "recount": ("Hint 1 (1st wrong answer)", "the columns beat on «दहाई» / «इकाई»"),
    "order": ("Hint 1 (1st wrong answer)", "the ones digits light on «इकाई», the tens digits on «दहाई»"),
    "h2": ("Hint 2 (2nd wrong answer)", "the numbers light; the tens and ones left beat on their words"),
    "h2_mental": ("Hint 2 (2nd wrong answer)", "the table comes in with the first number's blocks; the second number is taken away as it is said"),
    "reveal": ("Hint 3 (3rd wrong answer)", "the right card / the right keys lit (one key at a time), the others off"),
    "correct": ("correct answer", "the answer goes into the ?; confetti"),
    "eq": ("the equation read out", "after the correct line: each part lights as it is said"),
    "rm_tens_more": ("take-away mistake - 1st time", "a ten too many taken: the tens go back, this line"),
    "rm_guide": ("take-away mistake - 2nd time", "the tens go back; only the tens to take stay live, the others dimmed"),
    "rm_ones_more": ("take-away mistake (ones)", "not heard in this build: the ones lock once the right number is taken"),
    "rm_ones_less": ("take-away mistake (ones)", "not heard in this build: the tens are not tappable while ones are still to go"),
    "idle_build": ("inactivity reminder - building", "after 7s with no tap at the build step, again every 7s while idle"),
    "idle_rm": ("inactivity reminder - taking away", "after 7s with no tap at the take-away step, again every 7s while idle"),
    "idle_rm_now": ("take-away step said at once", "on a page with one take-away line: «अब इकाई हटाइए।» right after it, «अब दहाई हटाइए।» as the ones are done (tens pulse); then as the 7s reminder"),
    "cel": ("celebration", "the celebration screen: Swiftee jumps on «शाबाश», then says the line lip-synced"),
}


def page_rows(i, s):
    """the clips of one page in the order they are heard: [(audio id, (what, when))]"""
    sid, au, d, ty = s["id"], s["audio"], s["data"], s["type"]
    out = []
    def add(cid, key):
        if cid:
            out.append((cid, ROLE[key]))
    if ty == "STORY_SUB":
        for r in ("a", "b", "ask", "learn"):
            add(au.get(r), (sid, r))
    elif ty == "PV_SUB_DEMO":
        for r in ("head", "intro", "digits", "tens", "ones", "built", "sub", "rm_ones"):
            add(au.get(r), (sid, r))
        for n in (1, 2, 3):
            add("vo_num_%d" % n, (sid, "num"))
        for r in ("ones_left", "rm_tens"):
            add(au.get(r), (sid, r))
        add("vo_idle_rm_t", "idle_rm")
        for r in ("tens_left", "total", "story_a", "story_b", "story_c"):
            add(au.get(r), (sid, r))
    elif ty == "PV_SUB":
        b_t, b_o = d["b"] // 10, d["b"] % 10
        add(au.get("prompt"), "prompt")
        if d.get("build"):
            add(au.get("build"), "build"); add("vo_idle_build_t", "idle_build"); add("vo_idle_build_o", "idle_build")
        add(au.get("sub"), "sub")
        if d.get("steps"):
            add(au.get("rm_ones"), "rm_ones"); add(au.get("zero"), "zero"); add(au.get("rm_tens"), "rm_tens")
            if b_o: add("vo_idle_rm_o", "idle_rm")
            add("vo_idle_rm_t", "idle_rm")
        else:
            add(au.get("remove"), "remove"); add("vo_idle_rm_o", "idle_rm_now"); add("vo_idle_rm_t", "idle_rm_now")
        add(au.get("rm_tens_more"), "rm_tens_more"); add(au.get("rm_guide"), "rm_guide")
        add(au.get("ask"), "ask"); add(au.get("recount"), "recount"); add(au.get("h2"), "h2"); add(au.get("reveal"), "reveal")
        add(au.get("correct"), "correct"); add(au.get("eq"), "eq")
    elif ty == "MENTAL_SUB":
        for r, k in (("prompt", "prompt"), ("order", "order"), ("h2", "h2_mental"), ("reveal", "reveal"), ("correct", "correct")):
            add(au.get(r), k)
    elif ty == "CELEBRATION":
        add(au.get("prompt"), "cel")
    return out


# ---- the game order ----
slides = CARD["slides"]
order = [("Start screen", "Start screen", [("vo_landing", ("greeting", "the start screen: Swiftee's greeting; the picture comes in on «आज हम…»"))]),
         ("Transition 1", "before page 1", [("vo_pt_tutorial", ("transition line", "Swiftee rises talking; «चलिए, शुरू करें!» wipes in as it is said"))])]
gate_before = {"G1": ("Transition 2", "before page 3", "vo_pt_guided", "«चलिए, साथ में करें!»"),
               "P1": ("Transition 3", "before page 6", "vo_pt_practice", "«अब आपकी बारी!»")}
for i, s in enumerate(slides):
    if s["id"] in gate_before:
        g = gate_before[s["id"]]
        order.append((g[0], g[1], [(g[2], ("transition line", "Swiftee rises talking; " + g[3] + " wipes in as it is said"))]))
    page = "Celebration" if s["type"] == "CELEBRATION" else "Page %d" % (i + 1)
    screen = s["id"] + " · " + s["type"] + ("" if "a" not in s["data"] else "  (%d − %d)" % (s["data"]["a"], s["data"]["b"]))
    order.append((page, screen, page_rows(i, s)))

# where else each recording is heard
heard = {}
for page, screen, rows in order:
    for cid, _ in rows:
        heard.setdefault(cid, [])
        if page not in heard[cid]:
            heard[cid].append(page)

DUR = {cid: duration(os.path.join(VO_DIR, cid + ".ogg")) for cid in heard}
def status(cid):
    if not os.path.exists(os.path.join(VO_DIR, cid + ".ogg")):
        return "MISSING"
    if cid in COPIED: return "Copied (HI02H11, recorded voice)"
    rec = os.path.join(ROOT, "vo_source", "recorded")
    if os.path.isdir(rec) and any(os.path.splitext(f)[0] == cid for f in os.listdir(rec)): return "Recorded (2026-10-27)"
    return "TTS (Gemini, Kore)"

manifest = []
n = 0
for page, screen, rows in order:
    for k, (cid, (what, when)) in enumerate(rows, 1):
        n += 1
        also = ", ".join(p for p in heard[cid] if p != page)
        manifest.append(([n, page, screen, k, cid, TEXT.get(cid, ""), what, when, also, DUR[cid], status(cid)], status(cid)))

recordings = []
for page, screen, rows in order:
    for cid, _ in rows:
        if cid not in [r[0][1] for r in recordings]:
            recordings.append(([len(recordings) + 1, cid, cid + ".ogg", TEXT.get(cid, ""), ", ".join(heard[cid]), DUR[cid], status(cid)], status(cid)))
# in the card but never heard (kept for safety)
unheard = sorted(k for k in CARD["assets"]["audio"] if k.startswith("vo_") and k not in heard)

SFX = [("sfx_play", "Play button on the start screen - on the press", "Start"),
       ("sfx_next", "Next arrow (pages 1, 2) - on a real press", "1, 2"),
       ("sfx_gate_whoosh", "The three transition screens - as each opens (HI02H11's)", "transitions"),
       ("sfx_whoosh", "Start screen - the two children glide in", "Start"),
       ("sfx_correct", "Correct answer; the answer box filling", "2-10"),
       ("sfx_wrong", "Wrong answer / a wrong take-away move", "3-10"),
       ("sfx_confetti", "Confetti on a correct answer", "2-10"),
       ("sfx_nudge", "A button / block pulsing again when the child is idle", "2-10"),
       ("bgm", "Background music - from the play button on, at 30 %, dipped under the voice; paused while the tab is hidden", "all")]
sfx_rows = []
for sid, what, pages in SFX:
    f = os.path.join(SFX_DIR, sid + ".ogg")
    sfx_rows.append(([sid, "assets/sfx/" + sid + ".ogg", what, pages, duration(f) if os.path.exists(f) else "MISSING"], None))
f = os.path.join(VO_DIR, "sfx_celebrate.ogg")
sfx_rows.append((["sfx_celebrate", "assets/voiceover/sfx_celebrate.ogg", "Celebration screen - as it opens", "Celebration",
                  duration(f) if os.path.exists(f) else "MISSING"], None))

# ---- the workbook ----
thin = Side(style="thin", color="C9CED6"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill = PatternFill("solid", fgColor="1F4E79"); DEV = "Nirmala UI"
PAGE_FILL = {"Start screen": "EAF2FB", "Celebration": "FFF4D6"}


def sheet(ws, title, sub, head, widths, data, dev_col=None, wrap_cols=(), center=(), freeze="E5", fill=None):
    ws.append([title]); ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws.append([sub]); ws["A2"].font = Font(italic=True, size=10, color="555555")
    ws.append([])
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=4, column=c)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hfill; cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row, meta in data:
        ws.append(row); r = ws.max_row
        for c in range(1, len(head) + 1):
            cell = ws.cell(row=r, column=c); cell.border = box
            if fill: cell.fill = PatternFill("solid", fgColor=fill(row))
            cell.alignment = Alignment(vertical="center", wrap_text=(c in wrap_cols), horizontal="center" if c in center else None)
            if dev_col and c == dev_col:
                cell.font = Font(name=DEV, size=12, bold=True)
        if meta == "MISSING":
            ws.cell(row=r, column=len(head)).font = Font(bold=True, color="B03A2E")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    if freeze: ws.freeze_panes = freeze


def page_fill(row):
    p = str(row[1])
    if p.startswith("Transition"): return "EEF7EE"
    return PAGE_FILL.get(p, "FFFFFF" if int(row[0]) % 2 else "F7F9FC")


wb = Workbook()
ws = wb.active; ws.title = "VO Manifest"
sheet(ws, TITLE + " - lesson vo",
      "Start screen -> transitions and pages 1-10 in game order -> celebration. On each page: what plays as it opens, the steps, "
      "the reminders, the mistakes and the 1st / 2nd / 3rd wrong answer (Hint 1, 2, 3), then the correct answer.",
      ["S.No", "Page", "Screen", "#", "Audio ID", "Hindi script (spoken exactly)", "What it is", "When it plays", "Same recording also on",
       "Duration (s)", "Status"],
      [6, 13, 30, 4, 22, 50, 30, 62, 22, 10, 18], manifest, dev_col=6, wrap_cols=(3, 6, 7, 8, 9), center=(1, 4, 10), fill=page_fill)
sheet(wb.create_sheet("Recording list"), "Recording list - each recording once, in the order it is first heard",
      "%d recordings. Record each once; the file name is the Audio ID + .ogg, in assets/voiceover/." % len(recordings),
      ["S.No", "Audio ID", "File name", "Hindi script (spoken exactly)", "Plays on", "Duration (s)", "Status"],
      [6, 22, 24, 54, 44, 10, 22], recordings, dev_col=4, wrap_cols=(4, 5), center=(1, 6))
sheet(wb.create_sheet("Sound effects"), "Sound effects and music (supplied - not recorded)",
      "Only the files the lesson plays. The small interface sounds (taps, a block appearing / moving / popping, a block set "
      "down) are made in the browser (WebAudio) - no file.",
      ["Audio ID", "File name", "What it is / where it plays", "Page(s)", "Duration (s)"],
      [20, 34, 84, 12, 11], sfx_rows, wrap_cols=(3,), center=(5,), freeze=None)
sp = wb.create_sheet("Recording spec")
SPEC = [
    (TITLE + " - recording spec", None), (None, None),
    ("Scope", "The whole lesson - start screen, 3 transitions, 10 pages (T1, T2, G1-G3, P1-P5), celebration."),
    ("Order", "Game order. On each page: what plays as it opens, the steps, the 7s inactivity reminders, the take-away mistakes, "
              "then the 1st / 2nd / 3rd wrong answer (Hint 1, 2, 3) and the correct answer. A recording heard on more than one page "
              "is listed on each ('Same recording also on')."),
    ("Clips", "%d rows, %d distinct recordings (see 'Recording list')." % (len(manifest), len(recordings)) +
              (" In the card but not heard in this build (not listed): " + ", ".join(unheard) + "." if unheard else "")),
    ("Folder", "assets/voiceover/<Audio ID>.ogg - the file name is the Audio ID."),
    ("Voice", "Recorded voice (2026-10-27; sources in vo_source/recorded/, put in by tools/import_vo.py: ends trimmed, the lesson's "
              "lead / tail and loudness). Each clip was checked "
              "by transcription."),
    ("Format", "Ogg/Opus, 48 kHz mono, 32 kbps; 0.28 s before the voice, 0.36 s after; -18.5 LUFS (two-pass loudnorm)."),
    ("Register", "आप form throughout (the transition and celebration lines included)."),
    ("Script", "Record exactly the text in the Hindi script column - it is the build's own text. Numbers are spoken as Hindi words "
               "(0 = शून्य, 53 = तिरपन); a number written in digits is read as its word."),
    ("Timing", "Several pages light / move things on given words of a clip (e.g. «इसमें» lights the table, «मीरा» / «अमन» light a "
               "picture, «तेरह घटाए» sends the blocks to Aman). The times are measured from each clip's length: after new "
               "recordings run  python tools/lessons.py  so they are measured again from the new takes."),
    ("Celebration", "Swiftee's lip-sync on the celebration line is measured from vo_cel.ogg by tools/lessons.py (tools/make_lipsync.py) - "
                    "re-run it after a new take of vo_cel."),
    ("New takes", "Put the take in vo_source/recorded/ under its Audio ID (any audio format), then:  python tools/import_vo.py  "
                  "python tools/lessons.py  and  node tools/gen_asset_sizes.js  (bump the ?v= stamp in index.html so browsers fetch "
                  "the new audio)."),
    ("Regenerate", "python tools/gen_vo_manifest.py"),
]
for r, (a, b_) in enumerate(SPEC, 1):
    sp.cell(row=r, column=1, value=a); sp.cell(row=r, column=2, value=b_)
    sp.cell(row=r, column=1).font = Font(bold=True, size=14 if r == 1 else 11, color="1F4E79" if b_ is None else "000000")
    if b_:
        sp.cell(row=r, column=2).font = Font(name=DEV, size=11)
        sp.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        sp.cell(row=r, column=1).alignment = Alignment(vertical="top")
sp.column_dimensions["A"].width = 16; sp.column_dimensions["B"].width = 120

wb.save(OUT)
miss = [r[0][1] for r in recordings if r[1] == "MISSING"]
print("lesson_vo_manifest.xlsx: %d rows, %d recordings%s" % (len(manifest), len(recordings), ("; MISSING: " + ", ".join(miss)) if miss else ""))
if unheard:
    print("in the card, not heard:", ", ".join(unheard))
